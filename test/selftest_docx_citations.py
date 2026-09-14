"""docx 引文标记保护自检：上标/引用域引文编号在写回与润色中不被改写成 [1]。

覆盖：citation_key 规范化口径、写回侧（普通+修订模式）格式改写零改动、
真实改号仍生效且上标保留、last_rpr 泄漏回归（紧跟上标的普通文字不被
误染上标）、编辑器侧 protect_citations_in_text 还原逻辑。
（无 LLM 无网络，纯 python-docx + 本地写回内核。）
"""

from __future__ import annotations

import os
import sys
import tempfile
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

from src.core.docx_io import (  # noqa: E402
    _merge_para_text, citation_key, protect_citations_in_text, read_docx,
    write_docx,
)

RESULTS: list[tuple[bool, str]] = []


def check(name: str, cond: bool) -> None:
    RESULTS.append((bool(cond), name))
    print(f"[{'PASS' if cond else 'FAIL'}] {name}")


class _FakePara:
    """最小段落桩：_merge_para_text 只用 p._p。"""

    def __init__(self, p_el):
        self._p = p_el


def _build_docx(path: Path) -> None:
    """构造"适应性的演化研究¹表明先前工作存在不足。"——¹ 为上标 run。"""
    doc = Document()
    p = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    p.add_run("适应性的演化研究")
    sup = p.add_run("1")
    p.add_run("表明先前工作存在不足。")
    rpr = sup._r.get_or_add_rPr()
    rpr.append(rpr.makeelement(qn("w:vertAlign"), {qn("w:val"): "superscript"}))
    doc.save(path)


def _sup_texts(path: Path) -> list[str]:
    doc = Document(path)
    out = []
    for p in doc.paragraphs:
        for r in p.runs:
            rpr = r._r.find(qn("w:rPr"))
            if rpr is not None:
                va = rpr.find(qn("w:vertAlign"))
                if va is not None and va.get(qn("w:val")) == "superscript":
                    out.append(r.text)
    return out


# ============================================================
# 1. citation_key 规范化口径
# ============================================================
check("口径：[1] → 1", citation_key("[1]") == "1")
check("口径：6,7 原样", citation_key("6,7") == "6,7")
check("口径：中文括号+空格", citation_key("（10, 11）") == "10,11")
check("口径：[6-9] 区间", citation_key("[6-9]") == "6-9")
check("口径：年份四位数不是引文", citation_key("2020") is None)
check("口径：普通词语不是引文", citation_key("机制1") is None)
check("口径：空串安全", citation_key("") is None)

# ============================================================
# 2. 编辑器侧文本级保护
# ============================================================
# "研究 1 表明"：标记 "1" 在 [3,4)
toks = [(3, 4, "1")]
check("文本级：上标 1 被套上 [1] → 还原",
      protect_citations_in_text("研究 1 表明", "研究 [1] 表明", toks)
      == "研究 1 表明")
check("文本级：圆括号变体同样还原",
      protect_citations_in_text("研究 1 表明", "研究 (1) 表明", toks)
      == "研究 1 表明")
check("文本级：真实改号 1→2 不拦截",
      protect_citations_in_text("研究 1 表明", "研究 2 表明", toks)
      == "研究 2 表明")
check("文本级：非标记数字 3 不保护",
      protect_citations_in_text("研究 3 表明", "研究 [3] 表明", toks)
      == "研究 [3] 表明")
check("文本级：多数字标记 6,7",
      protect_citations_in_text("机制 6,7 报道", "机制 [6,7] 报道", [(3, 6, "6,7")])
      == "机制 6,7 报道")
check("文本级：括号裸化 [1]→1 还原",
      protect_citations_in_text("[1] 表示", "1 表示", [(0, 3, "[1]")])
      == "[1] 表示")
check("文本级：无标记时原样返回",
      protect_citations_in_text("研究 表明", "研究 [1] 表明", [])
      == "研究 [1] 表明")

# ============================================================
# 3. 写回保护（普通模式）
# ============================================================
tmp = Path(tempfile.mkdtemp(prefix="paperwb_cite_")) / "cite.docx"
_build_docx(tmp)
content = read_docx(tmp)
check("读入：识别出上标记文标记区间", content.citation_tokens[0] == [(8, 9, "1")])

write_docx(tmp, ["适应性的演化研究 [1] 表明先前工作存在局限。"])
check("写回：AI 套 [1] 后上标 run 原样保留", _sup_texts(tmp) == ["1"])
check("写回：正文文本仍是 1（非 [1]）",
      read_docx(tmp).paragraphs[0] == "适应性的演化研究1表明先前工作存在局限。")

# ============================================================
# 4. 写回保护（修订模式）：不再因域/格式改写整段退化
# ============================================================
_build_docx(tmp)
write_docx(tmp, ["适应性的演化研究 [1] 表明先前工作存在局限。"],
           track_changes=True)
check("修订写回：上标 run 原样保留", _sup_texts(tmp) == ["1"])
doc = Document(tmp)
check("修订写回：其余真实修改仍以 w:ins/w:del 记录",
      doc.paragraphs[0]._p.find(qn("w:ins")) is not None
      or doc.paragraphs[0]._p.find(qn("w:del")) is not None)

# ============================================================
# 5. 真实改号仍生效且继承上标
# ============================================================
_build_docx(tmp)
write_docx(tmp, ["适应性的演化研究2表明先前工作存在局限。"])
check("写回：改号 1→2 后上标保留", _sup_texts(tmp) == ["2"])

# ============================================================
# 6. last_rpr 泄漏回归：紧跟上标的普通文字重建时不被误染上标
# ============================================================
_build_docx(tmp)
doc = Document(tmp)
p = doc.paragraphs[0]
new_text = "适应性的演化研究1表明先前工作存在局限。"
_merge_para_text(p, new_text, None)  # 只有"不足→局限"真实变化
doc.save(tmp)
check("回归：无 rPr 的普通文字不继承上标", _sup_texts(tmp) == ["1"])

tmp.unlink()
print()
passed = sum(1 for ok, _ in RESULTS if ok)
failed = sum(1 for ok, _ in RESULTS if not ok)
print(f"{passed} passed, {failed} failed")
sys.exit(1 if failed else 0)
